from fastapi import FastAPI, HTTPException, Query, status, Depends, File, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import models
import observations
import shed
import movements
import logging
import threading
import os
import admin
from sqlalchemy import func, or_
from sqlalchemy.orm import Session, contains_eager
from typing import Annotated, Optional
from datetime import datetime
from math import ceil
from database import get_db, engine, ensure_zone_schema, SessionLocal, ensure_phone_unique_index, ensure_telegram_unique_index, ensure_inventory_schema
import pytz 
from dtos.itemResponseDTO import ItemResponseDTO
from dtos.deleteItemDTO import DeleteItemDTO, ResponseFakeDeleteDTO
import dtos.itemToCreateDTO as itemDTO
from historial import router
from auth import get_current_user, get_user_name_by_id, has_admin_access, router as auth_router
from notifications import NotificationService, enviar_mail_fallo_borrado
import zones
from seed_admin import seed_admin_from_env
from item_service import ItemServiceError, adjust_item_stock, create_item
from item_import import build_import_template, import_items_from_excel
from item_transfer import (
    ExportChecklistRequest,
    build_transfer_checklist,
    resolve_export_items,
    update_items_from_excel,
)
from whatsapp.router import router as whatsapp_router
from telegram.bot import router as telegram_router, start_telegram_bot
from item_images import router as item_images_router, delete_stored_image
import categories
import obras
import empleados
import existencias
import sysadmin
from item_categories import category_is_consumable, category_units_by_name, normalize_lookup, seed_categories
from pydantic import BaseModel
from observations import add_unit_observation
from unit_service import (
    apply_piece_profile,
    breaks_by_unit,
    sync_breakage,
    create_units_for_item,
    discard_unit,
    find_unit_by_code,
    change_item_counting,
    identify_current_stock,
    histories_by_unit,
    last_unit_movement,
    peek_codes,
    depot_zones_by_unit,
    pending_places_by_unit,
    reassign_unit_subcategory,
    set_unit_depot_zones,
    suggest_prefix,
)
from dotenv import load_dotenv

load_dotenv()

logging.basicConfig(level=logging.DEBUG)
logger = logging.getLogger(__name__)


app = FastAPI()

_raw_origins = os.getenv("ALLOWED_ORIGINS", "*").strip()
_allow_origins = (
    ["*"]
    if _raw_origins == "*"
    else [origin.strip() for origin in _raw_origins.split(",") if origin.strip()]
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=_allow_origins,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
    expose_headers=["*"]
)


app.include_router(router)
app.include_router(auth_router)
app.include_router(observations.router)
app.include_router(shed.router)
app.include_router(movements.router)
app.include_router(admin.router)
app.include_router(zones.router)
app.include_router(whatsapp_router)
app.include_router(telegram_router)
app.include_router(item_images_router)
app.include_router(categories.router)
app.include_router(obras.router)
app.include_router(empleados.router)
app.include_router(existencias.router)
app.include_router(sysadmin.router)

models.Base.metadata.create_all(bind=engine)
ensure_zone_schema()
ensure_inventory_schema()
_seed_db = SessionLocal()
try:
    seed_categories(_seed_db)
    from obras import seed_obras_from_history
    seed_obras_from_history(_seed_db)
finally:
    _seed_db.close()
try:
    from whatsapp.phone import canonicalize_stored_phones
    _phone_db = SessionLocal()
    try:
        canonicalize_stored_phones(_phone_db)
    finally:
        _phone_db.close()
except Exception:
    logger.exception("No se pudieron normalizar teléfonos de usuarios")
ensure_phone_unique_index()
ensure_telegram_unique_index()
seed_admin_from_env()

item_dependency = Annotated[Session, Depends(get_db)]

@app.on_event("startup")
def startup_event():
    if not hasattr(app, 'notification_thread'):
        app.notification_thread = threading.Thread(
            target=NotificationService.run_scheduler,
            daemon=True
        )
        app.notification_thread.start()
    if not getattr(app, "telegram_started", False):
        start_telegram_bot()
        app.telegram_started = True

TIMEZONE = pytz.timezone('America/Argentina/Buenos_Aires')

def now():
    """Devuelve la fecha/hora actual en la zona horaria de Buenos Aires"""
    return datetime.now(TIMEZONE)



DEFAULT_PAGE_SIZE = 50
MAX_PAGE_SIZE = 100

@app.get("/", response_model=dict)
def read_items(
    db: Session = Depends(get_db),
    name: Optional[str] = None,
    category: Optional[str] = None,
    shed_id: Optional[int] = None,
    zone_id: Optional[int] = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(DEFAULT_PAGE_SIZE, le=MAX_PAGE_SIZE)
):
    try:
        query = (
            db.query(models.Item)
            .outerjoin(models.Shed, models.Item.shed_id == models.Shed.id)
            .outerjoin(models.Zone, models.Item.zone_id == models.Zone.id)
            .options(contains_eager(models.Item.zone))
            .filter(models.Item.status == 1)
        )

        if name:
            needle = f"%{_fold_text(name)}%"
            piece_match = (
                db.query(models.ItemUnit.id)
                .filter(
                    models.ItemUnit.item_id == models.Item.id,
                    or_(
                        _fold_column(models.ItemUnit.code).like(needle),
                        _fold_column(func.coalesce(models.ItemUnit.name, "")).like(needle),
                    ),
                )
                .exists()
            )
            query = query.filter(or_(
                _fold_column(func.coalesce(models.Item.name, "")).like(needle),
                piece_match,
            ))
        if category:
            query = query.filter(models.Item.category.ilike(f"%{category}%"))
        if shed_id:
            query = query.filter(models.Item.shed_id == shed_id)
        if zone_id:
            query = query.filter(models.Item.zone_id == zone_id)

        total_records = query.count()
        total_pages = ceil(total_records / page_size)

        consumable_keys = {
            normalize_lookup(category.name)
            for category in db.query(models.Category).filter(models.Category.is_consumable == True).all()
        }
        unit_by_category = category_units_by_name(db)

        items = query.order_by(
                        models.Shed.name.asc().nullslast(),
                        models.Zone.name.asc().nullslast(),
                        models.Item.name.asc(),
                     ) \
                     .offset((page - 1) * page_size) \
                     .limit(page_size) \
                     .all()

        data = []
        for item in items:
            dto = ItemResponseDTO.model_validate(item)
            dto.zone_name = item.zone.name if item.zone else None
            dto.has_image = bool(item.image_filename)
            dto.track_units = bool(item.track_units)
            dto.inner_quantity = bool(item.inner_quantity)
            dto.is_consumable = normalize_lookup(item.category) in consumable_keys
            dto.unit = unit_by_category.get(normalize_lookup(item.category), "unidad")
            data.append(dto)

        return {
            "data": data,
            "pagination": {
                "total_records": total_records,
                "total_pages": total_pages,
                "current_page": page,
                "page_size": page_size,
                "has_next": page < total_pages,
                "has_previous": page > 1
            }
        }

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"Error al obtener items: {str(e)}"
        )

@app.get("/search")
def searchItems(name: str, db: item_dependency):
    needle = f"%{_fold_text(name)}%"
    rows = (
        db.query(models.Item.id, models.Item.name, models.Item.category)
        .filter(
            models.Item.status == 1,
            _fold_column(func.coalesce(models.Item.name, "")).like(needle),
        )
        .all()
    )
    return [{"id": row.id, "name": row.name, "category": row.category} for row in rows]


def getItemById(item_id: int, db: item_dependency):
    item = db.query(models.Item).filter(models.Item.id == item_id).first()
    if not item:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Item not found"
        )
    return item.name


@app.post("/")
def createItem(item: itemDTO.ItemCreateDTO, db: item_dependency):
    try:
        created = create_item(
            db,
            name=item.name,
            description=item.description,
            category=item.category,
            quantity=item.quantity,
            zone_id=item.zone_id,
            shed_id=item.shed_id,
            track_units=item.track_units,
            inner_quantity=item.inner_quantity,
            codes=item.codes,
            code_prefix=item.code_prefix,
        )
        return {
            "id": created.id,
            "name": created.name,
            "category": created.category,
            "track_units": bool(created.track_units),
            "inner_quantity": bool(created.inner_quantity),
            "code_prefix": created.code_prefix,
            "codes": list(getattr(created, "_created_codes", []) or []),
        }
    except ItemServiceError as e:
        raise HTTPException(status_code=e.status_code, detail=e.message)


@app.get("/items/import/template")
def download_items_import_template(db: item_dependency):
    content = build_import_template(db)
    return StreamingResponse(
        BytesIO(content),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": "attachment; filename=plantilla_carga_inventario.xlsx"
        },
    )


@app.post("/items/import")
async def import_items_excel(
    db: item_dependency,
    current_user: Annotated[dict, Depends(get_current_user)],
    file: UploadFile = File(...),
):
    filename = (file.filename or "").lower()
    if not filename.endswith(".xlsx"):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="El archivo debe ser un Excel (.xlsx)",
        )

    file_bytes = await file.read()
    if not file_bytes:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="El archivo está vacío",
        )

    try:
        return import_items_from_excel(db, file_bytes, current_user)
    except ItemServiceError as e:
        raise HTTPException(status_code=e.status_code, detail=e.message)
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))


@app.post("/items/export/traslado")
def export_transfer_checklist(
    payload: ExportChecklistRequest,
    db: item_dependency,
    current_user: Annotated[dict, Depends(get_current_user)],
):
    try:
        content = build_transfer_checklist(db, resolve_export_items(payload))
    except ItemServiceError as e:
        raise HTTPException(status_code=e.status_code, detail=e.message)

    filename = f"traslado_stock_{datetime.now(TIMEZONE).strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        BytesIO(content),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@app.post("/items/import/update")
async def import_items_update_excel(
    db: item_dependency,
    current_user: Annotated[dict, Depends(get_current_user)],
    file: UploadFile = File(...),
):
    filename = (file.filename or "").lower()
    if not filename.endswith(".xlsx"):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="El archivo debe ser un Excel (.xlsx)",
        )

    file_bytes = await file.read()
    if not file_bytes:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="El archivo está vacío",
        )

    try:
        return update_items_from_excel(db, file_bytes, current_user)
    except ItemServiceError as e:
        raise HTTPException(status_code=e.status_code, detail=e.message)
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))

@app.put("/items/by-id/{item_id}")
def update_item_by_id(
    item_id: int,
    item_update: itemDTO.ItemUpdateDTO,
    db: Session = Depends(get_db),
    current_user: Annotated[dict, Depends(get_current_user)] = None,
):
    item = db.query(models.Item).filter(models.Item.id == item_id).first()
    if not item:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Ítem con ID {item_id} no encontrado"
        )

    if item_update.quantity is None:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Debe especificar una cantidad"
        )

    if item_update.quantity < 0:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="La cantidad debe ser mayor a cero"
        )

    if item_update.contents and not has_admin_access(current_user):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Solo los administradores pueden cambiar la cantidad de un código",
        )

    from unit_service import add_inner_boxes, normalize_code, rename_unit, set_inner_contents

    try:
        for row in item_update.renames or []:
            unit = (
                db.query(models.ItemUnit)
                .filter(models.ItemUnit.id == row.id, models.ItemUnit.item_id == item.id)
                .first()
            )
            if not unit:
                raise ItemServiceError("La pieza no pertenece a esta subcategoría")
            rename_unit(unit, row.name)

        if item.inner_quantity:
            if item_update.action == itemDTO.ActionEnum.rest and (item_update.quantity or 0) > 0:
                raise ItemServiceError("Para sacar de una caja usá retirar")
            if item_update.contents:
                set_inner_contents(db, item, item_update.contents)
            pieces = item_update.piece_names or []
            if pieces:
                add_inner_boxes(db, item, pieces)
            elif (item_update.quantity or 0) > 0:
                raise ItemServiceError(
                    "Agregá una caja con su código, su nombre y cuántos hay adentro"
                )
            if not item_update.renames and not item_update.contents and not pieces:
                raise ItemServiceError("No hay cambios para guardar")
            db.commit()
            db.refresh(item)
            return {
                "id": item.id,
                "actualAmount": item.actualAmount,
                "totalAmount": item.totalAmount,
                "codes": list(getattr(item, "_created_codes", []) or []),
            }

        if item_update.quantity == 0:
            if not item_update.renames:
                raise ItemServiceError("No hay nombres para guardar")
            db.commit()
            db.refresh(item)
            return {
                "id": item.id,
                "actualAmount": item.actualAmount,
                "totalAmount": item.totalAmount,
                "codes": [],
            }

        quantity_change = 0

        if item_update.action == itemDTO.ActionEnum.add:
            quantity_change = item_update.quantity
        elif item_update.action == itemDTO.ActionEnum.rest:
            quantity_change = -item_update.quantity
        else:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Acción no válida"
            )

        codes = item_update.codes
        names_by_code = None
        if item.track_units and item_update.action == itemDTO.ActionEnum.add:
            pieces = item_update.piece_names or []
            if len(pieces) != item_update.quantity:
                raise ItemServiceError("Cada pieza nueva necesita un nombre")
            codes = []
            names_by_code = {}
            for piece in pieces:
                code = normalize_code(piece.code)
                codes.append(code)
                names_by_code[code] = piece.name

        updated = adjust_item_stock(
            db, item, quantity_change, codes=codes, names_by_code=names_by_code
        )
    except ItemServiceError as e:
        db.rollback()
        raise HTTPException(status_code=e.status_code, detail=e.message)

    return {
        "id": updated.id,
        "actualAmount": updated.actualAmount,
        "totalAmount": updated.totalAmount,
        "codes": list(getattr(updated, "_created_codes", []) or []),
    }

@app.put("/")
def updateItem(name: str, quantity: int, db: item_dependency):
    item = db.query(models.Item).filter(models.Item.name == name).first()
    if not item:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Item not found"
        )
    if quantity < 0:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Quantity cannot be negative"
        )
    item.totalAmount += quantity
    item.actualAmount += quantity
    db.commit()
    db.refresh(item)
    return item


@app.get("/items/{item_id}")
def get_item_details(
    item_id: int,
    db: item_dependency,
    current_user: Annotated[dict, Depends(get_current_user)]
):

    try:

        item = db.query(models.Item).filter(models.Item.id == item_id).first()

        if not item:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Item con ID {item_id} no encontrado"
            )

        is_deleted = item.status == 0

        deletion_history = None
        if is_deleted:
            deletion_history = db.query(models.DeletedItem)\
                .filter(models.DeletedItem.original_id == item_id)\
                .order_by(models.DeletedItem.deleted_at.desc())\
                .first()

        item_data = {
            "id": item.id,
            "name": item.name,
            "category": item.category,
            "description": item.description,
            "totalAmount": item.totalAmount,
            "actualAmount": item.actualAmount,
            "is_available": item.is_available,
            "shed_id": item.shed_id,
            "zone_id": item.zone_id,
            "status": item.status,
            "image_filename": item.image_filename,
            "has_image": bool(item.image_filename),
            "track_units": bool(item.track_units),
            "inner_quantity": bool(item.inner_quantity),
            "code_prefix": item.code_prefix,
        }
        item_data["is_consumable"] = category_is_consumable(db, item.category)
        item_data["unit"] = category_units_by_name(db).get(normalize_lookup(item.category), "unidad")

        response_data = {
            "item": item_data,
            "metadata": {
                "is_deleted": is_deleted,
                "deletion_info": {
                    "deletion_reason": deletion_history.deletion_reason if deletion_history else None,
                    "deleted_at": deletion_history.deleted_at if deletion_history else None
                } if is_deleted else None,
                "permissions": {
                    "can_edit": current_user["role"] in ["admin", "editor"],
                    "can_delete": current_user["role"] == "admin"
                }
            },
            "relations": {
                "observations_count": db.query(func.count(models.Observation.id)).filter(models.Observation.item_id == item.id).scalar() or 0,
                "movements_count": db.query(func.count(models.Movement.id)).filter(models.Movement.item_id == item.id).scalar() or 0,
                "last_movement": db.query(func.max(models.Movement.date)).filter(models.Movement.item_id == item.id).scalar(),
            }
        }

        return response_data



    except Exception as e:
        logger.error(f"Error obteniendo detalles del item {item_id}: {str(e)}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Error interno al obtener detalles del ítem"
        )



def _piece_place(unit: models.ItemUnit, item: models.Item | None, places) -> str:
    if unit.status == "consumida":
        return "Usada"
    if unit.status == "fuera_de_servicio":
        return "Fuera de servicio"
    if unit.status == "en_stock":
        zone = item.zone.name if item and item.zone else None
        return zone or "En depósito"
    names = [row["place"] for row in (places or []) if row.get("place")]
    return " · ".join(names) if names else "En obra"


def _fold_text(value):
    folded = value or ""
    for source, target in (
        ("á", "a"), ("Á", "a"), ("é", "e"), ("É", "e"),
        ("í", "i"), ("Í", "i"), ("ó", "o"), ("Ó", "o"),
        ("ú", "u"), ("Ú", "u"), ("ü", "u"), ("Ü", "u"),
        ("ñ", "n"), ("Ñ", "n"),
    ):
        folded = folded.replace(source, target)
    return folded.lower()


def _fold_column(column):
    folded = column
    for source, target in (
        ("á", "a"), ("Á", "a"), ("é", "e"), ("É", "e"),
        ("í", "i"), ("Í", "i"), ("ó", "o"), ("Ó", "o"),
        ("ú", "u"), ("Ú", "u"), ("ü", "u"), ("Ü", "u"),
        ("ñ", "n"), ("Ñ", "n"),
    ):
        folded = func.replace(folded, source, target)
    return func.lower(folded)


@app.get("/units")
def list_pieces(
    db: item_dependency,
    current_user: Annotated[dict, Depends(get_current_user)],
    q: str = "",
    broken: bool = False,
    without_image: bool = False,
    limit: int = Query(10, ge=1, le=500),
):
    text = (q or "").strip()
    broad = bool(text or broken or without_image)
    cap = 10 if not broad else min(limit, 500)
    query = (
        db.query(models.ItemUnit, models.Item)
        .join(models.Item, models.Item.id == models.ItemUnit.item_id)
        .outerjoin(models.Zone, models.Item.zone_id == models.Zone.id)
        .options(contains_eager(models.Item.zone))
        .filter(models.Item.status == 1, models.ItemUnit.status != "tirada")
    )
    if text:
        needle = f"%{_fold_text(text)}%"
        query = query.filter(or_(
            _fold_column(models.ItemUnit.code).like(needle),
            _fold_column(func.coalesce(models.ItemUnit.name, "")).like(needle),
            _fold_column(models.Item.name).like(needle),
        ))
    if broken:
        query = query.filter(models.ItemUnit.is_broken.is_(True))
    if without_image:
        query = query.filter(or_(
            models.ItemUnit.image_filename.is_(None),
            models.ItemUnit.image_filename == "",
        ))
    rows = query.order_by(models.ItemUnit.code.asc()).limit(cap).all()
    away_ids = [
        unit.id
        for unit, _item in rows
        if unit.status not in ("consumida", "fuera_de_servicio", "en_stock")
    ]
    places_by_unit = pending_places_by_unit(db, away_ids)
    data = []
    for unit, item in rows:
        data.append({
            "id": unit.id,
            "code": unit.code,
            "name": unit.name or "",
            "status": unit.status,
            "place": _piece_place(unit, item, places_by_unit.get(unit.id)),
            "is_broken": bool(unit.is_broken),
            "has_image": bool(unit.image_filename),
            "image_filename": unit.image_filename,
            "item_id": item.id,
            "item_name": item.name,
            "category": item.category,
        })
    return {"data": data}


def _pieces_export(rows, places_by_unit) -> BytesIO:
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Piezas"
    headers = [
        "Código",
        "Nombre",
        "Subcategoría",
        "Estado",
        "Dónde está",
        "Cantidad",
        "Qué le pasó",
        "Diagnóstico",
        "En reparación",
    ]
    labels = {
        "en_stock": "En depósito",
        "retirada": "En obra",
        "consumida": "Usada",
        "fuera_de_servicio": "Fuera de servicio",
    }
    header_fill = PatternFill("solid", fgColor="228BE6")
    header_font = Font(bold=True, color="FFFFFF")
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(1, col, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for row_index, (unit, item) in enumerate(rows, start=2):
        places = places_by_unit.get(unit.id) or []
        if item.inner_quantity:
            quantity = int(unit.quantity or 0)
            if quantity <= 0:
                quantity = sum(int(place.get("quantity") or 0) for place in places)
        else:
            quantity = 1
        values = [
            unit.code or "",
            unit.name or "",
            item.name or "",
            labels.get(unit.status, unit.status or ""),
            _piece_place(unit, item, places),
            quantity,
            unit.damage_note or "",
            unit.repair_note or "",
            "Sí" if unit.status == "fuera_de_servicio" else "No",
        ]
        for col, value in enumerate(values, start=1):
            sheet.cell(row_index, col, value)
    widths = [22, 36, 36, 22, 32, 12, 40, 40, 16]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


@app.get("/units/export")
def export_pieces(
    db: item_dependency,
    current_user: Annotated[dict, Depends(get_current_user)],
    q: str = "",
    broken: bool = False,
    without_image: bool = False,
):
    text = (q or "").strip()
    if not text and not broken and not without_image:
        raise HTTPException(status_code=400, detail="Elegí rotas, sin foto o un texto para exportar")
    query = (
        db.query(models.ItemUnit, models.Item)
        .join(models.Item, models.Item.id == models.ItemUnit.item_id)
        .outerjoin(models.Zone, models.Item.zone_id == models.Zone.id)
        .options(contains_eager(models.Item.zone))
        .filter(models.Item.status == 1, models.ItemUnit.status != "tirada")
    )
    if text:
        needle = f"%{_fold_text(text)}%"
        query = query.filter(or_(
            _fold_column(models.ItemUnit.code).like(needle),
            _fold_column(func.coalesce(models.ItemUnit.name, "")).like(needle),
            _fold_column(models.Item.name).like(needle),
        ))
    if broken:
        query = query.filter(models.ItemUnit.is_broken.is_(True))
    if without_image:
        query = query.filter(or_(
            models.ItemUnit.image_filename.is_(None),
            models.ItemUnit.image_filename == "",
        ))
    rows = query.order_by(models.ItemUnit.code.asc()).all()
    away_ids = [
        unit.id
        for unit, _item in rows
        if unit.status not in ("consumida", "fuera_de_servicio", "en_stock")
    ]
    places_by_unit = pending_places_by_unit(db, away_ids)
    buffer = _pieces_export(rows, places_by_unit)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="piezas.xlsx"'},
    )


@app.get("/units/suggest-prefix")
def suggest_unit_prefix(
    db: item_dependency,
    current_user: Annotated[dict, Depends(get_current_user)],
    name: str = Query(..., min_length=1),
):
    try:
        prefix = suggest_prefix(db, name)
    except ItemServiceError as e:
        raise HTTPException(status_code=e.status_code, detail=e.message)
    return {"prefix": prefix}


@app.get("/units/next-codes")
def next_unit_codes(
    db: item_dependency,
    current_user: Annotated[dict, Depends(get_current_user)],
    count: int = Query(1, ge=1, le=500),
    prefix: str = Query("K"),
):
    try:
        codes = peek_codes(db, count, prefix)
    except ItemServiceError as e:
        raise HTTPException(status_code=e.status_code, detail=e.message)
    return {"codes": codes}


@app.get("/units/by-code/{code}")
def get_unit_by_code(
    code: str,
    db: item_dependency,
    current_user: Annotated[dict, Depends(get_current_user)],
):
    try:
        unit = find_unit_by_code(db, code)
    except ItemServiceError as e:
        raise HTTPException(status_code=e.status_code, detail=e.message)
    if not unit:
        raise HTTPException(status_code=404, detail="No hay una pieza con ese código")

    item = db.query(models.Item).filter(models.Item.id == unit.item_id).first()
    history = last_unit_movement(db, unit)
    status_label = {
        "en_stock": "En depósito",
        "retirada": "En obra",
        "consumida": "Usada",
        "fuera_de_servicio": "Fuera de servicio",
    }.get(unit.status, unit.status)
    return {
        "id": unit.id,
        "code": unit.code,
        "name": unit.name,
        "is_broken": bool(unit.is_broken),
        "damage_note": unit.damage_note,
        "repair_note": unit.repair_note,
        "status": unit.status,
        "status_label": status_label,
        "created_at": unit.created_at.isoformat() if unit.created_at else None,
        "consumed_at": unit.consumed_at.isoformat() if unit.consumed_at else None,
        "item_id": item.id if item else unit.item_id,
        "item_name": item.name if item else None,
        "category": item.category if item else None,
        "shed_id": item.shed_id if item else None,
        "zone_id": item.zone_id if item else None,
        "zone_name": item.zone.name if item and item.zone else None,
        "last_action": history.action.value if history and history.action else None,
        "last_place": history.place if history else None,
        "last_person": history.personWhoTook if history else None,
        "last_date": history.date.isoformat() if history and history.date else None,
    }


@app.get("/items/{item_id}/units")
def list_item_units(
    item_id: int,
    db: item_dependency,
    current_user: Annotated[dict, Depends(get_current_user)],
    status_filter: Optional[str] = Query(None, alias="status"),
):
    item = db.query(models.Item).filter(models.Item.id == item_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Item not found")
    query = db.query(models.ItemUnit).filter(
        models.ItemUnit.item_id == item_id,
        models.ItemUnit.status != "tirada",
    )
    if status_filter and status_filter != "all":
        query = query.filter(models.ItemUnit.status == status_filter)
    elif not status_filter:
        query = query.filter(models.ItemUnit.status == "en_stock")
    units = query.order_by(models.ItemUnit.code.asc()).all()
    unit_ids = [unit.id for unit in units]
    places_by_unit = pending_places_by_unit(db, unit_ids)
    history_by_unit = histories_by_unit(db, unit_ids)
    depot_by_unit = depot_zones_by_unit(db, unit_ids)
    breaks = breaks_by_unit(db, unit_ids)
    labels = {
        "en_stock": "En depósito",
        "retirada": "En obra",
        "consumida": "Usada",
        "fuera_de_servicio": "Fuera de servicio",
    }
    payload = []
    for unit in units:
        places = places_by_unit.get(unit.id, [])
        payload.append({
            "id": unit.id,
            "code": unit.code,
            "status": unit.status,
            "status_label": labels.get(unit.status, unit.status),
            "quantity": int(unit.quantity or 0),
            "out_quantity": sum(row["quantity"] for row in places) if item.inner_quantity else 0,
            "out_places": places,
            "name": unit.name,
            "description": unit.description,
            "is_broken": bool(unit.is_broken),
            "damage_note": unit.damage_note,
            "repair_note": unit.repair_note,
            "has_image": bool(unit.image_filename),
            "image_filename": unit.image_filename,
            "created_at": unit.created_at.isoformat() if unit.created_at else None,
            "consumed_at": unit.consumed_at.isoformat() if unit.consumed_at else None,
            "history": history_by_unit.get(unit.id, []),
            "depot_zones": depot_by_unit.get(unit.id, []),
            "breaks": breaks.get(unit.id, []),
        })
    return {
        "track_units": bool(item.track_units),
        "inner_quantity": bool(item.inner_quantity),
        "is_consumable": category_is_consumable(db, item.category),
        "unit": category_units_by_name(db).get(normalize_lookup(item.category), "unidad"),
        "units": payload,
    }


class PieceCreateBody(BaseModel):
    code: Optional[str] = None
    name: str
    description: Optional[str] = None
    observation: Optional[str] = None
    is_broken: bool = False
    damage_note: Optional[str] = None
    repair_note: Optional[str] = None
    quantity: Optional[int] = None


class MoveSubcategoryBody(BaseModel):
    target_item_id: int
    code: Optional[str] = None


class ZonePlaceBody(BaseModel):
    zone_id: int
    quantity: int


class UnitZonesBody(BaseModel):
    places: list[ZonePlaceBody]


class DiscardBody(BaseModel):
    person: str
    authorized_by: str


class PieceUpdateBody(BaseModel):
    name: str
    description: Optional[str] = None
    is_broken: bool = False
    damage_note: Optional[str] = None
    repair_note: Optional[str] = None
    fix_note: Optional[str] = None


def _piece_response(unit) -> dict:
    return {
        "id": unit.id,
        "code": unit.code,
        "name": unit.name,
        "description": unit.description,
        "is_broken": bool(unit.is_broken),
        "damage_note": unit.damage_note,
        "repair_note": unit.repair_note,
        "quantity": int(unit.quantity or 0),
        "item_id": unit.item_id,
    }


@app.post("/items/{item_id}/pieces")
def create_piece(
    item_id: int,
    body: PieceCreateBody,
    db: item_dependency,
    current_user: Annotated[dict, Depends(get_current_user)],
):
    item = db.query(models.Item).filter(models.Item.id == item_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Item not found")
    if not item.track_units:
        raise HTTPException(status_code=400, detail="Esta subcategoría no identifica piezas")
    try:
        units = create_units_for_item(
            db,
            item,
            1,
            codes=[body.code] if body.code and body.code.strip() else None,
        )
        unit = units[0]
        apply_piece_profile(unit, body.name, body.is_broken, body.damage_note, body.repair_note, body.description)
        sync_breakage(db, unit, False)
        content = 1
        if item.inner_quantity:
            if body.quantity is None or body.quantity < 1:
                raise ItemServiceError("Indicá cuántos hay adentro")
            content = int(body.quantity)
            unit.quantity = content
        item.actualAmount = (item.actualAmount or 0) + content
        item.totalAmount = (item.totalAmount or 0) + content
        add_unit_observation(db, item.id, unit.id, body.observation, current_user)
        db.commit()
        db.refresh(unit)
    except ItemServiceError as e:
        db.rollback()
        raise HTTPException(status_code=e.status_code, detail=e.message)
    return _piece_response(unit)


@app.post("/items/{item_id}/units/{unit_id}/subcategory")
def move_piece_subcategory(
    item_id: int,
    unit_id: int,
    body: MoveSubcategoryBody,
    db: item_dependency,
    current_user: Annotated[dict, Depends(get_current_user)],
):
    source = db.query(models.Item).filter(models.Item.id == item_id).first()
    if not source:
        raise HTTPException(status_code=404, detail="Item not found")
    unit = db.query(models.ItemUnit).filter(models.ItemUnit.id == unit_id).first()
    if not unit or unit.item_id != source.id:
        raise HTTPException(status_code=404, detail="Pieza no encontrada")
    target = db.query(models.Item).filter(models.Item.id == body.target_item_id).first()
    if not target:
        raise HTTPException(status_code=404, detail="La subcategoría destino no existe")
    try:
        reassign_unit_subcategory(db, source, unit, target, body.code)
        db.commit()
        db.refresh(unit)
    except ItemServiceError as e:
        db.rollback()
        raise HTTPException(status_code=e.status_code, detail=e.message)
    remaining = db.query(models.ItemUnit).filter(models.ItemUnit.item_id == source.id).count()
    return {
        "unit": _piece_response(unit),
        "target_name": target.name,
        "source_empty": remaining == 0,
    }


@app.post("/units/{unit_id}/discard")
def discard_piece(
    unit_id: int,
    body: DiscardBody,
    db: item_dependency,
    current_user: Annotated[dict, Depends(get_current_user)],
):
    unit = db.query(models.ItemUnit).filter(models.ItemUnit.id == unit_id).first()
    if not unit:
        raise HTTPException(status_code=404, detail="Pieza no encontrada")
    item = db.query(models.Item).filter(models.Item.id == unit.item_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Item not found")
    try:
        discard_unit(
            db,
            item,
            unit,
            body.person,
            body.authorized_by,
            current_user["user_id"],
            get_user_name_by_id(db, current_user["user_id"]),
        )
        db.commit()
        db.refresh(unit)
    except ItemServiceError as e:
        db.rollback()
        raise HTTPException(status_code=e.status_code, detail=e.message)
    return {"code": unit.code}


@app.put("/units/{unit_id}/zones")
def update_piece_zones(
    unit_id: int,
    body: UnitZonesBody,
    db: item_dependency,
    current_user: Annotated[dict, Depends(get_current_user)],
):
    unit = db.query(models.ItemUnit).filter(models.ItemUnit.id == unit_id).first()
    if not unit:
        raise HTTPException(status_code=404, detail="Pieza no encontrada")
    item = db.query(models.Item).filter(models.Item.id == unit.item_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Item not found")
    try:
        places = set_unit_depot_zones(db, item, unit, body.places)
        db.commit()
    except ItemServiceError as e:
        db.rollback()
        raise HTTPException(status_code=e.status_code, detail=e.message)
    return {"depot_zones": places}


@app.put("/units/{unit_id}")
def update_piece(
    unit_id: int,
    body: PieceUpdateBody,
    db: item_dependency,
    current_user: Annotated[dict, Depends(get_current_user)],
):
    unit = db.query(models.ItemUnit).filter(models.ItemUnit.id == unit_id).first()
    if not unit:
        raise HTTPException(status_code=404, detail="Pieza no encontrada")
    try:
        was_broken = bool(unit.is_broken)
        previous_damage = unit.damage_note
        previous_repair = unit.repair_note
        apply_piece_profile(unit, body.name, body.is_broken, body.damage_note, body.repair_note, body.description)
        sync_breakage(db, unit, was_broken, previous_damage, previous_repair, body.fix_note)
        db.commit()
        db.refresh(unit)
    except ItemServiceError as e:
        db.rollback()
        raise HTTPException(status_code=e.status_code, detail=e.message)
    return _piece_response(unit)


class IdentifyUnitsBody(BaseModel):
    prefix: str


class CountingBody(BaseModel):
    mode: str
    prefix: Optional[str] = None


def _counting_payload(item, units) -> dict:
    return {
        "id": item.id,
        "name": item.name,
        "track_units": bool(item.track_units),
        "inner_quantity": bool(item.inner_quantity),
        "code_prefix": item.code_prefix,
        "actualAmount": item.actualAmount,
        "totalAmount": item.totalAmount,
        "codes": [unit.code for unit in units],
    }


@app.put("/items/{item_id}/counting")
def update_item_counting(
    item_id: int,
    body: CountingBody,
    db: item_dependency,
    current_user: Annotated[dict, Depends(get_current_user)],
):
    if not has_admin_access(current_user):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Solo los administradores pueden cambiar cómo se cuenta",
        )
    item = db.query(models.Item).filter(models.Item.id == item_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Item not found")
    try:
        units = change_item_counting(db, item, body.mode, body.prefix)
    except ItemServiceError as e:
        raise HTTPException(status_code=e.status_code, detail=e.message)
    return _counting_payload(item, units)


@app.post("/items/{item_id}/identify")
def identify_item_units(
    item_id: int,
    body: IdentifyUnitsBody,
    db: item_dependency,
    current_user: Annotated[dict, Depends(get_current_user)],
):
    item = db.query(models.Item).filter(models.Item.id == item_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="Item not found")
    try:
        units = identify_current_stock(db, item, body.prefix)
    except ItemServiceError as e:
        raise HTTPException(status_code=e.status_code, detail=e.message)
    return {"codes": [unit.code for unit in units], "track_units": True}


@app.delete("/")
def delete_product(
    item_delete: DeleteItemDTO,
    db: item_dependency,
    current_user: Annotated[dict, Depends(get_current_user)]
):
    item = db.query(models.Item).filter(models.Item.id == item_delete.item_id).first()
    item_name = item.name if item else "desconocido"
    can_delete = item.actualAmount == item.totalAmount

    if not has_admin_access(current_user):
        fake_dto = ResponseFakeDeleteDTO(
            item_id=item_delete.item_id,
            description=item_delete.description,
            date=item_delete.date,
            username=current_user["username"]
        )
        enviar_mail_fallo_borrado(fake_dto, item_name)
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="No tenés permisos para borrar este ítem."
        )

    if not item:
        raise HTTPException(status_code=404, detail="Item no encontrado")

    if not can_delete:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Estas queriendo eliminar un producto cuando en este momento hay parte de ese producto en una obra."
        )

    # Registrar el borrado
    deleted_item = models.DeletedItem(
        item_id=item.id,
        name=item.name,
        description=item.description,
        category=item.category,
        status=item.status,
        deletion_reason=item_delete.description,
        deleted_at=now(),
        # original_id=item.id
    )
    db.add(deleted_item)

    # Eliminar todas las observaciones asociadas al item
    db.query(models.Observation).filter(
        models.Observation.item_id == item.id
    ).delete()

    delete_stored_image(item)

    # Marcar el item como borrado y cambiar su nombre
    item.status = 0
    item.name = f"{item.name}__DELETED_{datetime.utcnow().strftime('%Y%m%d%H%M%S')}"

    db.commit()

    return {"message": f"Item {item_delete.item_id} borrado exitosamente"}


@app.get("/deleted-items", response_model=dict)
def get_deleted_items(
    name: Optional[str] = Query(None),
    category: Optional[str] = Query(None),
    from_date: Optional[datetime] = Query(None, alias="from"),
    to_date: Optional[datetime] = Query(None, alias="to"),
    page: int = Query(1, ge=1),
    page_size: int = Query(DEFAULT_PAGE_SIZE, le=MAX_PAGE_SIZE),
    db: Session = Depends(get_db)
):
    try:
        query = db.query(models.DeletedItem)

        if name:
            query = query.filter(models.DeletedItem.name.ilike(f"%{name}%"))

        if category:
            query = query.filter(models.DeletedItem.category.ilike(f"%{category}%"))

        if from_date:
            query = query.filter(models.DeletedItem.deleted_at >= from_date)

        if to_date:
            query = query.filter(models.DeletedItem.deleted_at <= to_date)

        total_records = query.count()
        total_pages = ceil(total_records / page_size)

        deleted_items = (
            query.order_by(models.DeletedItem.deleted_at.desc())
            .offset((page - 1) * page_size)
            .limit(page_size)
            .all()
        )

        items_data = [
            {
                "id": item.id,
                "item_id": item.item_id,
                "name": item.name,
                "category": item.category,
                "description": item.description,
                "deletion_reason": item.deletion_reason,
                "deleted_at": item.deleted_at
            }
            for item in deleted_items
        ]

        return {
            "data": items_data,  # Ahora tiene la misma estructura que antes
            "pagination": {
                "total_records": total_records,
                "total_pages": total_pages,
                "current_page": page,
                "page_size": page_size,
                "has_next": page < total_pages,
                "has_previous": page > 1,
            },
        }

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"Error al obtener ítems eliminados: {str(e)}"
        )